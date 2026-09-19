from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import cast

import pytest
from boa_oi import portfolio_api
from boa_oi.api import application_for
from boa_oi.models.entities import (
    AuditLog,
    Customer,
    Opportunity,
    OpportunityAction,
    PortfolioAssignment,
    PropensityScoreRecord,
    RelationshipManager,
)
from boa_oi.platform import Principal, current_principal
from boa_oi.technical.ids import deterministic_uuid
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Table, create_engine
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool


@pytest.fixture(autouse=True)
def clean_dependency_overrides():
    yield
    for service in ("portfolio-service", "customer-service"):
        application_for(service).dependency_overrides.clear()


def principal(role: str, *, managers: tuple[str, ...] = (), branches: tuple[str, ...] = ()):
    return Principal(
        subject=f"subject-{role.lower()}",
        username=role.lower(),
        roles={role},
        scopes=set(),
        client_id="boa-sme-spa",
        relationship_manager_ids=managers,
        branch_ids=branches,
        customer_scopes=("assigned" if role == "RELATIONSHIP_MANAGER" else "branch",),
    )


def factory(*, with_assignments: bool = True):
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    tables = [
        RelationshipManager.__mapper__.local_table,
        Customer.__mapper__.local_table,
    ]
    if with_assignments:
        tables.extend(
            [
                PortfolioAssignment.__mapper__.local_table,
                PropensityScoreRecord.__mapper__.local_table,
                Opportunity.__mapper__.local_table,
                OpportunityAction.__mapper__.local_table,
                AuditLog.__mapper__.local_table,
            ]
        )
    with engine.connect() as connection:
        schemas = (
            ("customer", "ml", "opportunity", "action", "audit")
            if with_assignments
            else ("customer",)
        )
        for schema in schemas:
            connection.exec_driver_sql(f"ATTACH DATABASE ':memory:' AS '{schema}'")
        for table in tables:
            cast(Table, table).create(connection)
    return sessionmaker(bind=engine, expire_on_commit=False)


def client_for(service: str, session_factory, identity: Principal) -> TestClient:
    app = application_for(service)
    app.state.session_factory = session_factory
    app.dependency_overrides[current_principal] = lambda: identity
    return TestClient(app)


def seed_reassigned_customer(
    session_factory,
) -> tuple[Customer, RelationshipManager, RelationshipManager]:
    old_manager = RelationshipManager(
        id=deterministic_uuid("rm", "old"),
        subject_id="rm-old",
        display_name="Ancien CC",
        branch_code="BR-01",
        active=True,
    )
    new_manager = RelationshipManager(
        id=deterministic_uuid("rm", "new"),
        subject_id="rm-new",
        display_name="Nouveau CC",
        branch_code="BR-02",
        active=True,
    )
    customer = Customer(
        id=deterministic_uuid("customer", "SME-00999"),
        customer_ref="SME-00999",
        legal_name="Entreprise Réaffectée",
        sector_code="SERVICES",
        segment_code="MEDIUM",
        scenario_code="NORMAL_CUSTOMER",
        incorporated_on=date(2018, 1, 1),
        status="ACTIVE",
        # Deliberately stale: active scopes must no longer depend on this legacy pointer.
        rm_id=old_manager.id,
    )
    now = datetime.now(timezone.utc)
    with session_factory.begin() as session:
        session.add_all([old_manager, new_manager, customer])
        session.add_all(
            [
                PortfolioAssignment(
                    id=deterministic_uuid("assignment", "old"),
                    customer_id=customer.id,
                    relationship_manager_id=old_manager.id,
                    branch_code="BR-01",
                    valid_from=now - timedelta(days=60),
                    valid_to=now - timedelta(days=1),
                    actor="portfolio-admin",
                    reason="Historique avant réaffectation",
                ),
                PortfolioAssignment(
                    id=deterministic_uuid("assignment", "new"),
                    customer_id=customer.id,
                    relationship_manager_id=new_manager.id,
                    branch_code="BR-02",
                    valid_from=now - timedelta(days=1),
                    valid_to=None,
                    actor="portfolio-admin",
                    reason="Réaffectation vers la nouvelle agence",
                ),
            ]
        )
    return customer, old_manager, new_manager


def test_active_assignment_drives_customer_scope_and_serialized_branch(monkeypatch):
    monkeypatch.setenv("BOA_ALLOW_NON_POSTGRES_TEST_DB", "true")
    session_factory = factory()
    seed_reassigned_customer(session_factory)

    old_cc = client_for(
        "customer-service",
        session_factory,
        principal("RELATIONSHIP_MANAGER", managers=("rm-old",), branches=("BR-01",)),
    )
    assert old_cc.get("/internal/v1/customers/SME-00999").status_code == 404

    new_cc = client_for(
        "customer-service",
        session_factory,
        principal("RELATIONSHIP_MANAGER", managers=("rm-new",), branches=("BR-02",)),
    )
    response = new_cc.get("/internal/v1/customers/SME-00999")
    assert response.status_code == 200
    assert response.json()["relationshipManagerId"] == "rm-new"
    assert response.json()["branchId"] == "BR-02"


def test_active_assignment_drives_portfolio_dashboard_and_branch_scope(monkeypatch):
    monkeypatch.setenv("BOA_ALLOW_NON_POSTGRES_TEST_DB", "true")
    session_factory = factory()
    seed_reassigned_customer(session_factory)

    old_branch = client_for(
        "portfolio-service",
        session_factory,
        principal("BRANCH_MANAGER", branches=("BR-01",)),
    )
    assert (
        old_branch.get("/internal/v1/dashboards/branch").json()["kpis"]["portfolioCustomers"] == 0
    )

    new_cc = client_for(
        "portfolio-service",
        session_factory,
        principal("RELATIONSHIP_MANAGER", managers=("rm-new",), branches=("BR-02",)),
    )
    payload = new_cc.get("/internal/v1/dashboards/me").json()
    assert payload["scope"]["relationshipManagerId"] == "rm-new"
    assert payload["scope"]["branchId"] == "BR-02"
    assert [item["customerId"] for item in payload["portfolio"]] == ["SME-00999"]
    assert payload["portfolio"][0]["branchId"] == "BR-02"


def test_only_one_active_assignment_is_allowed_per_customer():
    session_factory = factory()
    customer, _old_manager, new_manager = seed_reassigned_customer(session_factory)

    with pytest.raises(IntegrityError), session_factory.begin() as session:
        session.add(
            PortfolioAssignment(
                id=deterministic_uuid("assignment", "duplicate"),
                customer_id=customer.id,
                relationship_manager_id=new_manager.id,
                branch_code="BR-02",
                valid_from=datetime.now(timezone.utc),
                valid_to=None,
                actor="unit-test",
                reason="Duplicate active assignment must fail",
            )
        )


def test_legacy_sqlite_fixtures_without_assignment_table_keep_working(monkeypatch):
    monkeypatch.setenv("BOA_ALLOW_NON_POSTGRES_TEST_DB", "true")
    session_factory = factory(with_assignments=False)
    manager = RelationshipManager(
        id=deterministic_uuid("rm", "legacy"),
        subject_id="rm-legacy",
        display_name="CC Fixture Historique",
        branch_code="BR-01",
        active=True,
    )
    customer = Customer(
        id=deterministic_uuid("customer", "SME-LEGACY"),
        customer_ref="SME-LEGACY",
        legal_name="Fixture SQLite historique",
        sector_code="SERVICES",
        segment_code="SMALL",
        scenario_code="NORMAL_CUSTOMER",
        incorporated_on=date(2020, 1, 1),
        status="ACTIVE",
        rm_id=manager.id,
    )
    with session_factory.begin() as session:
        session.add_all([manager, customer])

    client = client_for(
        "customer-service",
        session_factory,
        principal("RELATIONSHIP_MANAGER", managers=("rm-legacy",), branches=("BR-01",)),
    )
    response = client.get("/internal/v1/customers/SME-LEGACY")
    assert response.status_code == 200
    assert response.json()["relationshipManagerId"] == "rm-legacy"
    assert response.json()["branchId"] == "BR-01"


def test_assignment_validity_metadata_is_persisted():
    session_factory = factory()
    customer, old_manager, _new_manager = seed_reassigned_customer(session_factory)
    with session_factory() as session:
        history = list(
            session.query(PortfolioAssignment)
            .filter(PortfolioAssignment.customer_id == customer.id)
            .order_by(PortfolioAssignment.valid_from)
        )
    assert len(history) == 2
    assert history[0].relationship_manager_id == old_manager.id
    assert history[0].valid_to is not None
    assert history[1].valid_to is None
    assert history[1].actor == "portfolio-admin"
    assert history[1].reason == "Réaffectation vers la nouvelle agence"


def test_portfolio_export_uses_active_assignment_and_hides_other_branch(monkeypatch):
    monkeypatch.setenv("BOA_ALLOW_NON_POSTGRES_TEST_DB", "true")
    session_factory = factory()
    seed_reassigned_customer(session_factory)

    new_cc = client_for(
        "portfolio-service",
        session_factory,
        principal("RELATIONSHIP_MANAGER", managers=("rm-new",), branches=("BR-02",)),
    )
    response = new_cc.get(
        "/internal/v1/exports/portfolio.xlsx",
        headers={"X-Correlation-ID": "corr-portfolio-export"},
    )
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    rows = list(workbook["Portefeuille PME"].iter_rows(values_only=True))
    assert [row[0] for row in rows[1:]] == ["SME-00999"]
    assert rows[1][4:7] == ("rm-new", "Nouveau CC", "Casablanca Sidi Maârouf")

    old_branch = client_for(
        "portfolio-service",
        session_factory,
        principal("BRANCH_MANAGER", branches=("BR-01",)),
    )
    denied = old_branch.get("/internal/v1/exports/portfolio.xlsx?relationshipManagerId=rm-new")
    assert denied.status_code == 404


def test_portfolio_export_rejects_volume_before_building_payload(monkeypatch):
    monkeypatch.setenv("BOA_ALLOW_NON_POSTGRES_TEST_DB", "true")
    monkeypatch.setattr(portfolio_api, "MAX_EXPORT_ROWS", 0)
    session_factory = factory()
    seed_reassigned_customer(session_factory)
    client = client_for(
        "portfolio-service",
        session_factory,
        principal("RELATIONSHIP_MANAGER", managers=("rm-new",), branches=("BR-02",)),
    )

    response = client.get("/internal/v1/exports/portfolio.xlsx")

    assert response.status_code == 413
    assert response.json()["code"] == "EXPORT_LIMIT_EXCEEDED"
