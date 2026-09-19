from __future__ import annotations

from io import BytesIO

import pytest
from boa_oi import export_xlsx
from openpyxl import load_workbook


def test_workbook_contains_data_metadata_and_no_formula() -> None:
    artifact = export_xlsx.build_workbook(
        filename="opportunites.xlsx",
        sheet_name="Opportunités",
        columns=export_xlsx.opportunity_columns(),
        rows=[
            {
                "opportunityId": "OPP-001",
                "customerId": "SME-001",
                "customerName": '=HYPERLINK("https://example.test")',
                "opportunityType": "TRADE_FINANCE",
                "status": "OPEN",
                "confidence": 0.8,
                "priorityScore": 72,
                "why": ["Encaissements en hausse"],
                "recommendedProducts": [{"name": "Trade Finance"}],
            }
        ],
        metadata={"exportType": "OPPORTUNITIES", "correlationId": "corr-xlsx"},
    )

    assert artifact.filename == "opportunites.xlsx"
    assert artifact.row_count == 1
    assert len(artifact.sha256) == 64
    workbook = load_workbook(BytesIO(artifact.content), read_only=True, data_only=False)
    assert workbook.sheetnames == ["Opportunités", "Métadonnées"]
    rows = list(workbook["Opportunités"].iter_rows(values_only=True))
    assert rows[0][0:3] == ("Identifiant opportunité", "Identifiant PME", "PME")
    assert rows[1][0:2] == ("OPP-001", "SME-001")
    assert rows[1][2].startswith("'=")
    assert "Trade Finance" in rows[1]
    metadata = dict(workbook["Métadonnées"].iter_rows(min_row=2, values_only=True))
    assert metadata["schemaVersion"] == "1.0"
    assert metadata["rowCount"] == 1
    assert metadata["correlationId"] == "corr-xlsx"
    assert "aucune décision de crédit" in metadata["disclaimer"]


def test_workbook_rejects_rows_above_the_pilot_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(export_xlsx, "MAX_EXPORT_ROWS", 1)

    with pytest.raises(ValueError, match="1-row pilot limit"):
        export_xlsx.build_workbook(
            filename="portfolio.xlsx",
            sheet_name="Portefeuille PME",
            columns=export_xlsx.portfolio_columns(),
            rows=[{"customerId": "SME-1"}, {"customerId": "SME-2"}],
            metadata={},
        )
