from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import date, datetime, timezone
from io import BytesIO
from typing import Any, Callable, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_SCHEMA_VERSION = "1.0"
MAX_EXPORT_ROWS = 10_000
DISCLAIMER = (
    "Pilotage commercial uniquement. Les signaux et priorités ne constituent aucune "
    "décision de crédit."
)


@dataclass(frozen=True)
class XlsxArtifact:
    content: bytes
    filename: str
    sha256: str
    row_count: int


Column = tuple[str, Callable[[Mapping[str, Any]], Any]]


def _safe_text(value: str) -> str:
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def safe_cell(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, str):
        return _safe_text(value)
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray)):
        rendered = "; ".join(str(safe_cell(item) or "") for item in value)
        return _safe_text(rendered)
    if isinstance(value, Mapping):
        rendered = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
        return _safe_text(rendered)
    return _safe_text(str(value))


def _header_cell(sheet: Any, value: str) -> Any:
    cell = WriteOnlyCell(sheet, value=value)
    cell.font = Font(color="FFFFFF", bold=True)
    cell.fill = PatternFill(fill_type="solid", fgColor="0B2B4C")
    return cell


def build_workbook(
    *,
    filename: str,
    sheet_name: str,
    columns: Sequence[Column],
    rows: Iterable[Mapping[str, Any]],
    metadata: Mapping[str, Any],
) -> XlsxArtifact:
    materialized = list(rows)
    if len(materialized) > MAX_EXPORT_ROWS:
        raise ValueError(f"Export exceeds the {MAX_EXPORT_ROWS}-row pilot limit.")

    workbook = Workbook(write_only=True)
    data_sheet = workbook.create_sheet(title=sheet_name)
    data_sheet.freeze_panes = "A2"
    data_sheet.append([_header_cell(data_sheet, label) for label, _extract in columns])
    for row in materialized:
        data_sheet.append([safe_cell(extract(row)) for _label, extract in columns])

    metadata_sheet = workbook.create_sheet(title="Métadonnées")
    metadata_sheet.append(
        [_header_cell(metadata_sheet, "Propriété"), _header_cell(metadata_sheet, "Valeur")]
    )
    complete_metadata = {
        "schemaVersion": EXPORT_SCHEMA_VERSION,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "rowCount": len(materialized),
        "disclaimer": DISCLAIMER,
        **metadata,
    }
    for key, value in complete_metadata.items():
        metadata_sheet.append([key, safe_cell(value)])

    output = BytesIO()
    workbook.save(output)
    content = output.getvalue()
    return XlsxArtifact(
        content=content,
        filename=filename,
        sha256=hashlib.sha256(content).hexdigest(),
        row_count=len(materialized),
    )


def opportunity_columns() -> tuple[Column, ...]:
    def products(row: Mapping[str, Any]) -> list[str]:
        values = row.get("recommendedProducts") or []
        return [
            str(item.get("name") or item.get("productId") or "")
            if isinstance(item, Mapping)
            else str(item)
            for item in values
        ]

    return (
        ("Identifiant opportunité", lambda row: row.get("opportunityId")),
        ("Identifiant PME", lambda row: row.get("customerId")),
        ("PME", lambda row: row.get("customerName") or row.get("legalName")),
        ("Type d'opportunité", lambda row: row.get("opportunityType")),
        ("Statut", lambda row: row.get("status")),
        ("Confiance", lambda row: row.get("confidence")),
        ("Niveau de confiance", lambda row: row.get("confidenceLevel")),
        ("Score de priorité", lambda row: row.get("priorityScore")),
        ("Niveau de priorité", lambda row: row.get("priorityLevel")),
        ("Horizon", lambda row: row.get("horizon")),
        ("Pourquoi", lambda row: row.get("why")),
        ("Action recommandée", lambda row: row.get("what")),
        ("Quand", lambda row: row.get("when")),
        ("Produits recommandés", products),
        ("Détectée le", lambda row: row.get("generatedAt")),
        ("Mise à jour statut", lambda row: row.get("statusUpdatedAt")),
        ("Expire le", lambda row: row.get("expiresAt")),
        ("Dernière action", lambda row: row.get("lastActionAt")),
        ("Version moteur", lambda row: row.get("engineVersion")),
        ("Version règle", lambda row: row.get("ruleVersion")),
        ("Politique de scoring", lambda row: row.get("scoringPolicyId")),
        ("Version politique", lambda row: row.get("scoringPolicyVersion")),
        ("Mode de calcul", lambda row: row.get("fallbackMode")),
    )


def portfolio_columns() -> tuple[Column, ...]:
    return (
        ("Identifiant PME", lambda row: row.get("customerId")),
        ("PME", lambda row: row.get("customerName")),
        ("Secteur", lambda row: row.get("industry")),
        ("Segment", lambda row: row.get("segment")),
        ("Identifiant CC", lambda row: row.get("relationshipManagerId")),
        ("Chargé de clientèle", lambda row: row.get("relationshipManagerName")),
        ("Agence", lambda row: row.get("branchName") or row.get("branchId")),
        ("Score de propension POC", lambda row: row.get("propensityScore")),
        ("Score de priorité", lambda row: row.get("combinedPriorityScore")),
        ("Niveau de priorité", lambda row: row.get("priorityLevel")),
        ("Motif de priorité", lambda row: row.get("priorityReason")),
        ("Opportunités ouvertes", lambda row: len(row.get("openOpportunities") or [])),
        (
            "Identifiants opportunités",
            lambda row: [
                item.get("opportunityId")
                for item in row.get("openOpportunities") or []
                if isinstance(item, Mapping)
            ],
        ),
        ("Actions à venir", lambda row: len(row.get("nextActions") or [])),
        (
            "Prochaine échéance",
            lambda row: min(
                (
                    str(item["dueAt"])
                    for item in row.get("nextActions") or []
                    if isinstance(item, Mapping) and item.get("dueAt")
                ),
                default=None,
            ),
        ),
    )


__all__ = [
    "DISCLAIMER",
    "EXPORT_SCHEMA_VERSION",
    "MAX_EXPORT_ROWS",
    "XLSX_MIME",
    "XlsxArtifact",
    "build_workbook",
    "opportunity_columns",
    "portfolio_columns",
    "safe_cell",
]
